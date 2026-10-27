"""Read and validate a filled-in history / reference workbook.

Rows are returned as plain dicts keyed by column key. Every problem is reported as
an ``ImportIssue`` (row, column, message); nothing is guessed or silently fixed.
Storing the rows in SQLite is Phase 1 work.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from teklif.library import Library
from teklif.reporting.excel_templates import DATA_SHEET, header
from teklif.storage.history_schema import Column, Workbook, list_values

Value = str | int | float | date | None


@dataclass(frozen=True)
class ImportIssue:
    row: int
    column: str
    message: str


@dataclass
class ImportResult:
    rows: list[dict[str, Value]] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.issues


def _coerce(col: Column, raw: Any, allowed: set[str] | None) -> Value:
    """Convert one cell to the column type; raise ValueError with a Turkish message."""
    if col.kind == "text":
        return str(raw).strip()
    if col.kind == "list":
        text = str(raw).strip()
        if allowed is not None and text not in allowed:
            raise ValueError(f"'{text}' listede yok")
        return text
    if col.kind == "int":
        if isinstance(raw, bool) or not isinstance(raw, int | float) or float(raw) != int(raw):
            raise ValueError(f"tam sayı bekleniyordu: {raw!r}")
        if raw < 0:
            raise ValueError("negatif olamaz")
        return int(raw)
    if col.kind == "float":
        if isinstance(raw, bool) or not isinstance(raw, int | float):
            raise ValueError(f"sayı bekleniyordu: {raw!r}")
        return float(raw)
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    raise ValueError(f"tarih bekleniyordu: {raw!r}")


def read_workbook(path: Path | str, spec: Workbook, lib: Library) -> ImportResult:
    result = ImportResult()
    wb = load_workbook(path, data_only=True, read_only=True)
    if DATA_SHEET not in wb.sheetnames:
        result.issues.append(ImportIssue(0, "", f"'{DATA_SHEET}' sayfası bulunamadı"))
        return result
    rows = wb[DATA_SHEET].iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    except StopIteration:
        result.issues.append(ImportIssue(1, "", "başlık satırı boş"))
        return result

    index: dict[str, int] = {}
    for col in spec.columns:
        for candidate in (header(col), col.key):
            if candidate in headers:
                index[col.key] = headers.index(candidate)
                break
        else:
            # Optional columns may be missing (e.g. a file made from an older template).
            if col.required:
                result.issues.append(ImportIssue(1, col.key, "kolon bulunamadı"))
    if result.issues:
        return result

    allowed = {
        c.key: {v for v, _ in list_values(lib, c.source)}
        for c in spec.columns
        if c.source is not None
    }
    for row_no, values in enumerate(rows, start=2):
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        record: dict[str, Value] = {}
        for col in spec.columns:
            i = index.get(col.key, -1)
            raw = values[i] if 0 <= i < len(values) else None
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                if col.required:
                    result.issues.append(ImportIssue(row_no, col.key, "zorunlu alan boş"))
                record[col.key] = None
                continue
            try:
                record[col.key] = _coerce(col, raw, allowed.get(col.key))
            except ValueError as exc:
                result.issues.append(ImportIssue(row_no, col.key, str(exc)))
                record[col.key] = None
        result.rows.append(record)
    wb.close()
    return result
