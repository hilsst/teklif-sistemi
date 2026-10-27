"""Build the empty upload workbooks with dropdowns from the column specs.

Each workbook has three sheets:
- ``Veri``: one header row; dropdowns and numeric/date validation on every data row.
- ``Aciklama``: column guide (description, unit, required, allowed values).
- ``Listeler``: allowed ids and labels for every dropdown, read from the library.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook as XlWorkbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from teklif.library import Library
from teklif.storage.history_schema import WORKBOOKS, Column, ListSource, Workbook, list_values

DATA_SHEET = "Veri"
GUIDE_SHEET = "Aciklama"
LISTS_SHEET = "Listeler"
DATA_ROWS = 2000

_REQUIRED_FILL = PatternFill("solid", fgColor="F4B183")
_OPTIONAL_FILL = PatternFill("solid", fgColor="D9E1F2")
_HEADER_FONT = Font(bold=True)
_KIND_LABELS = {
    "text": "metin",
    "int": "tam sayı",
    "float": "sayı",
    "date": "tarih",
    "list": "liste",
}


def header(col: Column) -> str:
    """Header text shown in Excel: key plus unit, e.g. ``kalip_alani [m2]``."""
    return f"{col.key} [{col.unit}]" if col.unit else col.key


def _write_lists(ws: Worksheet, lib: Library, sources: list[ListSource]) -> dict[str, str]:
    """Write each list in two columns (id, label); return source -> range formula."""
    ranges: dict[str, str] = {}
    for i, source in enumerate(sources):
        id_col = get_column_letter(2 * i + 1)
        label_col = get_column_letter(2 * i + 2)
        ws[f"{id_col}1"] = source
        ws[f"{label_col}1"] = "açıklama"
        ws[f"{id_col}1"].font = _HEADER_FONT
        values = list_values(lib, source)
        for r, (vid, label) in enumerate(values, start=2):
            ws[f"{id_col}{r}"] = vid
            ws[f"{label_col}{r}"] = label
        ws.column_dimensions[id_col].width = 22
        ws.column_dimensions[label_col].width = 36
        last = len(values) + 1
        ranges[source] = f"{LISTS_SHEET}!${id_col}$2:${id_col}${last}"
    return ranges


def _validation(col: Column, list_ranges: dict[str, str]) -> DataValidation | None:
    dv: DataValidation | None
    if col.kind == "list":
        assert col.source is not None
        dv = DataValidation(type="list", formula1=list_ranges[col.source])
        dv.error = "Lütfen listeden seçin (Listeler sayfası)."
    elif col.kind == "int":
        dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0")
        dv.error = "0 veya daha büyük bir tam sayı girin."
    elif col.kind == "float":
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="-1000000000")
        dv.error = "Sayı girin (ondalık ayırıcı Excel ayarınıza göre)."
    elif col.kind == "date":
        dv = DataValidation(type="date", operator="greaterThan", formula1="1")
        dv.error = "Tarih girin (GG.AA.YYYY)."
    else:
        return None
    dv.allow_blank = True
    dv.showErrorMessage = True
    dv.errorTitle = "Geçersiz değer"
    return dv


def build_workbook(spec: Workbook, lib: Library) -> XlWorkbook:
    wb = XlWorkbook()
    data = wb.active
    assert data is not None
    data.title = DATA_SHEET
    guide = wb.create_sheet(GUIDE_SHEET)
    lists = wb.create_sheet(LISTS_SHEET)

    sources: list[ListSource] = []
    for col in spec.columns:
        if col.source is not None and col.source not in sources:
            sources.append(col.source)
    list_ranges = _write_lists(lists, lib, sources)

    for idx, col in enumerate(spec.columns, start=1):
        letter = get_column_letter(idx)
        cell = data.cell(row=1, column=idx, value=header(col))
        cell.font = _HEADER_FONT
        cell.fill = _REQUIRED_FILL if col.required else _OPTIONAL_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        note = col.desc + (f"\nBirim: {col.unit}" if col.unit else "")
        note += "\nZorunlu" if col.required else ""
        cell.comment = Comment(note, "teklif")
        data.column_dimensions[letter].width = max(14, min(32, len(header(col)) + 2))
        dv = _validation(col, list_ranges)
        if dv is not None:
            dv.add(f"{letter}2:{letter}{DATA_ROWS + 1}")
            data.add_data_validation(dv)
        if col.kind == "date":
            for r in range(2, DATA_ROWS + 2):
                data.cell(row=r, column=idx).number_format = "DD.MM.YYYY"
    data.freeze_panes = "B2"
    data.row_dimensions[1].height = 32

    guide.append([spec.title])
    guide["A1"].font = Font(bold=True, size=14)
    guide.append([spec.row_meaning])
    guide.append(["Turuncu başlık = zorunlu, mavi = isteğe bağlı. Başlıkları değiştirmeyin."])
    guide.append([])
    guide.append(["kolon", "açıklama", "birim", "tip", "zorunlu", "geçerli değerler"])
    for c in guide[5]:
        c.font = _HEADER_FONT
    for col in spec.columns:
        allowed = ""
        if col.source is not None:
            allowed = ", ".join(v for v, _ in list_values(lib, col.source))
        guide.append(
            [
                header(col),
                col.desc,
                col.unit,
                _KIND_LABELS[col.kind],
                "evet" if col.required else "",
                allowed,
            ]
        )
    for letter, width in zip("ABCDEF", (34, 60, 18, 10, 9, 60), strict=True):
        guide.column_dimensions[letter].width = width

    return wb


def write_templates(lib: Library, out_dir: Path | str) -> list[Path]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for spec in WORKBOOKS.values():
        path = out / f"{spec.name}.xlsx"
        build_workbook(spec, lib).save(path)
        paths.append(path)
    return paths
