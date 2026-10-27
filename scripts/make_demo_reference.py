"""Build ``examples/demo_referans_teklifler.xlsx``: a DUMMY reference set for the demo.

"Actual" costs are the demo-library result multiplied by fixed factors, so the
regression report shows both passing and failing jobs. Not real data.

Usage: uv run python scripts/make_demo_reference.py  (after make_demo_library.py)
"""

from __future__ import annotations

import tempfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from teklif.costing.engine import compute_quote
from teklif.costing.regression import row_to_input
from teklif.library import load_library
from teklif.reporting.excel_templates import DATA_SHEET, write_templates
from teklif.storage.history_import import Value
from teklif.storage.history_schema import REFERENCE_QUOTES

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "examples" / "demo_referans_teklifler.xlsx"

# (part, material, plies, ply area total, part area, trim, holes, tool area, tool perim,
#  tool x, tool y, complexity, qty, batch, actual/computed factor)
JOBS = [
    (
        "DEMO-KAPAK",
        "PP_CF_WOVEN_120",
        8,
        9.6,
        1.0,
        4.2,
        12,
        1.8,
        5.6,
        1600,
        1100,
        "orta",
        20,
        5,
        0.95,
    ),
    (
        "DEMO-BRAKET",
        "PP_CF_UD_120",
        16,
        3.2,
        0.2,
        1.9,
        6,
        0.5,
        2.9,
        800,
        600,
        "yuksek",
        50,
        10,
        1.08,
    ),
    (
        "DEMO-PANEL",
        "PP_CF_WOVEN_120",
        6,
        13.8,
        2.3,
        6.4,
        0,
        3.0,
        7.4,
        2300,
        1500,
        "dusuk",
        10,
        2,
        0.88,
    ),
    (
        "DEMO-KORUYUCU",
        "PP_GF_WOVEN_120",
        4,
        6.0,
        1.5,
        5.0,
        4,
        2.2,
        6.2,
        1900,
        1300,
        "dusuk",
        40,
        8,
        1.12,
    ),
    (
        "DEMO-KANAL",
        "PP_CF_UD_120",
        12,
        7.2,
        0.6,
        3.8,
        20,
        1.2,
        4.6,
        1500,
        900,
        "yuksek",
        15,
        3,
        1.25,
    ),
    (
        "DEMO-KAPI",
        "PP_CF_WOVEN_120",
        10,
        18.0,
        1.8,
        6.0,
        30,
        2.8,
        7.0,
        2100,
        1400,
        "orta",
        8,
        2,
        0.80,
    ),
]


def main() -> int:
    lib = load_library(REPO / "examples" / "demo_library")
    rows: list[dict[str, Value]] = []
    for i, (
        part,
        mat,
        plies,
        ply_area,
        area,
        trim,
        holes,
        tarea,
        tper,
        tx,
        ty,
        cx,
        qty,
        batch,
        factor,
    ) in enumerate(JOBS, start=1):
        row: dict[str, Value] = {
            "teklif_no": f"DEMO-{i:03d}",
            "tarih": date(2026, 1, i),
            "musteri": "Demo müşteri",
            "parca_no": part,
            "proses": "autoclave_prepreg",
            "kur_sicakligi": 120.0,
            "malzeme": mat,
            "siparis_adedi": qty,
            "parti_adedi": batch,
            "ply_sayisi": plies,
            "toplam_ply_alani": ply_area,
            "parca_alani": area,
            "trim_cevresi": trim,
            "delik_sayisi": holes,
            "kalip_alani": tarea,
            "kalip_cevresi": tper,
            "kalip_zarf_x": float(tx),
            "kalip_zarf_y": float(ty),
            "ilk_parca_sira_no": 1,
            "karmasiklik": cx,
            "para_birimi": "EUR",
            "fiili_birim_maliyet": 0.0,
        }
        result = compute_quote(row_to_input(row, lib), lib)
        cost, price = result.value("cost.unit"), result.value("price.unit")
        assert cost is not None and price is not None
        row["fiili_birim_maliyet"] = round(cost * factor, 2)
        row["teklif_birim_fiyat"] = round(price * 1.05, 2)
        rows.append(row)

    with tempfile.TemporaryDirectory() as tmp:
        path = next(p for p in write_templates(lib, tmp) if p.stem == "referans_teklifler")
        wb = load_workbook(path)
        ws = wb[DATA_SHEET]
        for r, row in enumerate(rows, start=2):
            for c, col in enumerate(REFERENCE_QUOTES.columns, start=1):
                if col.key in row:
                    ws.cell(row=r, column=c, value=row[col.key])
        wb.save(OUT)
    print(f"Wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
