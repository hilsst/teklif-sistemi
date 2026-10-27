from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from teklif.cli import main
from teklif.common.models import Override
from teklif.costing.engine import compute_quote
from teklif.costing.inputs import QuoteInput
from teklif.costing.regression import run_regression
from teklif.library import Library
from teklif.reporting.excel_templates import DATA_SHEET, write_templates
from teklif.reporting.quote_report import render_html, write_excel, write_pdf
from teklif.storage.db import Store
from teklif.storage.history_import import Value
from teklif.storage.history_schema import REFERENCE_QUOTES

from .conftest import EXAMPLE_INPUT


def test_store_round_trip_with_override_log(
    tmp_path: Path, filled_lib: Library, example_input: QuoteInput
) -> None:
    inp = example_input.model_copy(
        update={"overrides": {"labor.serim": Override(value=50.0, reason="ölçülen süre")}}
    )
    result = compute_quote(inp, filled_lib)
    store = Store(tmp_path / "t.sqlite")
    first = store.save_quote(result, filled_lib)
    second = store.save_quote(result, filled_lib)
    assert (first.revision, second.revision) == (1, 2)
    loaded = store.get(first.id).result()
    assert loaded.value("price.order_total") == pytest.approx(result.value("price.order_total"))
    assert loaded.get("labor.serim").override is not None
    log = store.overrides(first.id)
    assert [(o.line_key, o.value, o.reason) for o in log] == [("labor.serim", 50.0, "ölçülen süre")]
    assert log[0].computed == pytest.approx(result.get("labor.serim").result)
    assert len(first.library_hash) == 64
    store.add_actual(first.id, 2900.0, labor_hours=4.5, note="ilk parti")
    assert [a.unit_cost for a in store.actuals(first.id)] == [2900.0]
    with pytest.raises(KeyError):
        store.add_actual(999, 1.0)


def test_excel_and_pdf_reports(
    tmp_path: Path, filled_lib: Library, example_input: QuoteInput
) -> None:
    result = compute_quote(example_input, filled_lib)
    wb = load_workbook(write_excel(result, tmp_path / "q.xlsx"))
    assert wb.sheetnames == ["Ozet", "Kirilim"]
    rows = list(wb["Kirilim"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(result.lines)
    assert rows[0][0] == "learning.factor"
    html = render_html(result)
    assert "Kalem kalem kırılım" in html and "time.serim" in html
    pytest.importorskip("weasyprint")
    pdf = write_pdf(result, tmp_path / "q.pdf")
    assert pdf.read_bytes().startswith(b"%PDF")


def _reference_file(tmp_path: Path, lib: Library, rows: list[dict[str, Value]]) -> Path:
    path = next(p for p in write_templates(lib, tmp_path) if p.stem == "referans_teklifler")
    wb = load_workbook(path)
    ws = wb[DATA_SHEET]
    for r, row in enumerate(rows, start=2):
        for c, col in enumerate(REFERENCE_QUOTES.columns, start=1):
            if col.key in row:
                ws.cell(row=r, column=c, value=row[col.key])
    wb.save(path)
    return path


REF_ROW: dict[str, Value] = {
    "teklif_no": "R-1",
    "tarih": date(2025, 5, 1),
    "parca_no": "P-9",
    "proses": "autoclave_prepreg",
    "kur_sicakligi": 120,
    "malzeme": "PP_CF_UD_120",
    "siparis_adedi": 20,
    "parti_adedi": 5,
    "ply_sayisi": 10,
    "toplam_ply_alani": 11.0,
    "parca_alani": 1.0,
    "trim_cevresi": 4.2,
    "delik_sayisi": 12,
    "kalip_alani": 1.8,
    "kalip_cevresi": 5.6,
    "kalip_zarf_x": 1600,
    "kalip_zarf_y": 1100,
    "karmasiklik": "orta",
    "para_birimi": "EUR",
    "teklif_birim_fiyat": 3500.0,
}


def test_regression_report(tmp_path: Path, filled_lib: Library) -> None:
    computed_ref = (
        run_regression(
            _reference_file(tmp_path / "a", filled_lib, [dict(REF_ROW, fiili_birim_maliyet=1.0)]),
            filled_lib,
        )
        .rows[0]
        .computed
    )
    assert computed_ref is not None
    rows = [
        dict(REF_ROW, teklif_no="R-ok", fiili_birim_maliyet=computed_ref * 1.10),
        dict(REF_ROW, teklif_no="R-far", fiili_birim_maliyet=computed_ref * 1.5),
        dict(REF_ROW, teklif_no="R-try", para_birimi="TRY", fiili_birim_maliyet=computed_ref * 40),
        dict(REF_ROW, teklif_no="R-eksik", kalip_alani=None, fiili_birim_maliyet=100.0),
        dict(REF_ROW, teklif_no="R-sicak", kur_sicakligi=150, fiili_birim_maliyet=100.0),
    ]
    report = run_regression(_reference_file(tmp_path / "b", filled_lib, rows), filled_lib)
    status = {r.quote_no: r.status for r in report.rows}
    assert status == {
        "R-ok": "ok",
        "R-far": "out_of_tolerance",
        "R-try": "ok",
        "R-eksik": "skipped",
        "R-sicak": "skipped",
    }
    by_no = {r.quote_no: r for r in report.rows}
    assert by_no["R-try"].deviation == pytest.approx(0.0)
    assert "kalip_alani" in by_no["R-eksik"].reason
    assert report.tolerance == pytest.approx(0.15)


def test_regression_with_todo_library_reports_manual_review(tmp_path: Path, lib: Library) -> None:
    path = _reference_file(tmp_path, lib, [dict(REF_ROW, fiili_birim_maliyet=100.0)])
    report = run_regression(path, lib)
    assert [r.status for r in report.rows] == ["manual_review"]
    assert report.compared == []


def test_cli_quote(tmp_path: Path, filled_lib_dir: Path) -> None:
    db = tmp_path / "q.sqlite"
    args = [
        "--library",
        str(filled_lib_dir),
        "quote",
        str(EXAMPLE_INPUT),
        "--xlsx",
        str(tmp_path / "q.xlsx"),
        "--db",
        str(db),
    ]
    assert main(args) == 0
    assert (tmp_path / "q.xlsx").exists()
    assert Store(db).list_quotes()[0].quote_no == "T-2026-001"
    bad = tmp_path / "bad.yaml"
    bad.write_text(
        EXAMPLE_INPUT.read_text(encoding="utf-8").replace("complexity: orta", "complexity: cok"),
        encoding="utf-8",
    )
    assert main(["--library", str(filled_lib_dir), "quote", str(bad)]) == 1


def test_streamlit_app_smoke(
    tmp_path: Path, filled_lib_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest

    monkeypatch.setenv("TEKLIF_DB", str(tmp_path / "app.sqlite"))
    app = Path(__file__).resolve().parents[1] / "src/teklif/app/streamlit_app.py"
    at = AppTest.from_file(str(app), default_timeout=60)
    at.run()
    at.sidebar.text_input[0].set_value(str(filled_lib_dir)).run()
    assert not at.exception
    metrics = {m.label: m.value for m in at.metric}
    assert metrics["Birim maliyet [EUR]"] != "—"
    labels = [b.label for b in at.button]
    at.button[labels.index("Teklifi kaydet")].click().run()
    assert not at.exception
    assert [s.value for s in at.success] == ["Kaydedildi: T-2026-001 rev 1"]
