from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from teklif.cli import main
from teklif.library import Library
from teklif.reporting.excel_templates import (
    DATA_SHEET,
    GUIDE_SHEET,
    LISTS_SHEET,
    header,
    write_templates,
)
from teklif.storage.history_import import Value, read_workbook
from teklif.storage.history_schema import LABOR_HISTORY, TOOL_HISTORY, WORKBOOKS


@pytest.fixture
def templates(lib: Library, tmp_path: Path) -> dict[str, Path]:
    return {p.stem: p for p in write_templates(lib, tmp_path)}


def test_all_workbooks_written(templates: dict[str, Path]) -> None:
    assert set(templates) == {"kalip_gecmis", "iscilik_gecmis", "referans_teklifler"}


@pytest.mark.parametrize("name", sorted(WORKBOOKS))
def test_headers_sheets_and_dropdowns(templates: dict[str, Path], name: str) -> None:
    spec = WORKBOOKS[name]
    wb = load_workbook(templates[name])
    assert wb.sheetnames == [DATA_SHEET, GUIDE_SHEET, LISTS_SHEET]
    ws = wb[DATA_SHEET]
    assert [c.value for c in ws[1]] == [header(c) for c in spec.columns]
    list_cols = [c for c in spec.columns if c.kind == "list"]
    list_dvs = [dv for dv in ws.data_validations.dataValidation if dv.type == "list"]
    assert len(list_dvs) == len(list_cols)
    assert all(dv.formula1.startswith("Listeler!$") for dv in list_dvs)


def test_lists_come_from_library(templates: dict[str, Path], lib: Library) -> None:
    ws = load_workbook(templates["iscilik_gecmis"])[LISTS_SHEET]
    values = {c.value for col in ws.iter_cols() for c in col}
    assert {w.id for w in lib.work_centers.work_centers} <= values
    assert {"EUR", "TRY", "USD"} & values == set()  # currencies not used in this workbook
    assert "torbalama" in values


def _fill(path: Path, row: dict[str, Value]) -> None:
    spec = WORKBOOKS[path.stem]
    wb = load_workbook(path)
    ws = wb[DATA_SHEET]
    for i, col in enumerate(spec.columns, start=1):
        if col.key in row:
            ws.cell(row=2, column=i, value=row[col.key])
    wb.save(path)


LABOR_ROW: dict[str, Value] = {
    "is_emri_no": "IE-001",
    "tarih": date(2026, 9, 1),
    "parca_no": "P-100",
    "proses": "autoclave_prepreg",
    "operasyon": "serim",
    "is_merkezi": "TEMIZ_ODA",
    "parti_adedi": 4,
    "ply_sayisi": 12,
    "toplam_ply_alani": 6.5,
    "fiili_iscilik": 960,
}


def test_valid_upload_round_trip(templates: dict[str, Path], lib: Library) -> None:
    path = templates["iscilik_gecmis"]
    _fill(path, LABOR_ROW)
    result = read_workbook(path, LABOR_HISTORY, lib)
    assert result.ok, result.issues
    assert len(result.rows) == 1
    row = result.rows[0]
    assert row["tarih"] == date(2026, 9, 1)
    assert row["parti_adedi"] == 4
    assert row["fiili_iscilik"] == pytest.approx(960.0)
    assert row["karmasiklik"] is None


def test_invalid_upload_reports_every_issue(templates: dict[str, Path], lib: Library) -> None:
    path = templates["iscilik_gecmis"]
    bad = dict(LABOR_ROW, is_merkezi="FIRIN_9", parti_adedi=2.5, fiili_iscilik=None)
    _fill(path, bad)
    result = read_workbook(path, LABOR_HISTORY, lib)
    problems = {(i.column, i.message.split(":")[0]) for i in result.issues}
    assert ("is_merkezi", "'FIRIN_9' listede yok") in problems
    assert ("parti_adedi", "tam sayı bekleniyordu") in problems
    assert ("fiili_iscilik", "zorunlu alan boş") in problems


def test_missing_column_is_reported(templates: dict[str, Path], lib: Library) -> None:
    path = templates["kalip_gecmis"]
    wb = load_workbook(path)
    wb[DATA_SHEET].delete_cols(1)
    wb.save(path)
    result = read_workbook(path, TOOL_HISTORY, lib)
    assert [(i.column, i.message) for i in result.issues] == [("kalip_no", "kolon bulunamadı")]


def test_cli_commands(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["validate"]) == 0
    assert main(["templates", "--out", str(tmp_path)]) == 0
    out = tmp_path / "todo.md"
    assert main(["todo", "--out", str(out)]) == 0
    assert out.read_text(encoding="utf-8").startswith("# Firma verisi")
    assert main(["check-upload", "kalip_gecmis", str(tmp_path / "kalip_gecmis.xlsx")]) == 0
    assert "0 satır okundu" in capsys.readouterr().out
